import json

import pymupdf
import pytest
from pptx import Presentation


def make_pdf(path, pages):
    doc = pymupdf.open()
    for text in pages:
        page = doc.new_page()
        page.insert_text((72, 72), text, fontsize=12)
    doc.save(path)


def make_pptx(path, slides):
    prs = Presentation()
    for title, body, notes in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
        slide.notes_slide.notes_text_frame.text = notes
    prs.save(path)


@pytest.fixture
def course_dir(tmp_path):
    """A small course: a PDF, a folder of notes and code, a notebook, a slide deck, and extras."""
    d = tmp_path / "course"
    (d / "lecture02").mkdir(parents=True)
    make_pdf(
        d / "MIT6_006F11_lec01.pdf",
        [
            "Introduction\nAlgorithms solve problems.\n1",
            "Peak Finding\nA peak is no smaller than its neighbors.",
            "MIT OpenCourseWare\nFor information about citing these materials or our Terms of Use, visit ...",
        ],
    )
    (d / "lecture02" / "notes.md").write_text("Intro text.\n\n# Graphs\n\nVertices and edges.\n\n## BFS\n\nLevels.\n")
    (d / "lecture02" / "bfs.py").write_text("def bfs(graph, s):\n    ...\n")
    nb = {
        "cells": [
            {"cell_type": "markdown", "source": ["# Sorting\n", "Insertion sort."]},
            {"cell_type": "code", "source": "sorted([3, 1, 2])"},
            {"cell_type": "markdown", "source": "## Merge sort\nDivide and conquer."},
        ]
    }
    (d / "lec03.ipynb").write_text(json.dumps(nb))
    make_pptx(d / "Lecture 4 - Hashing.pptx", [("Hashing", "Hash tables", "Say hello."), ("Chaining", "Lists", "")])
    (d / "syllabus.html").write_text(
        "<html><head><title>Syllabus</title></head><body><nav>Menu</nav>"
        "<main><h1>Syllabus</h1><p>Weekly problem sets.</p><h2>Grading</h2><p>Exams.</p></main></body></html>"
    )
    make_pdf(d / "tutorial3.pdf", ["Tutorial three"])
    (d / ".hidden.md").write_text("secret")
    return d


@pytest.fixture
def ocw_dir(tmp_path):
    """A small OpenCourseWare export: two chapters listed in a table, the full notes linked from a
    paragraph, a problem set, and a scanned page."""
    d = tmp_path / "ocw"
    (d / "static_resources").mkdir(parents=True)
    (d / "data.json").write_text(
        json.dumps({"site_short_id": "18.014-fall-2010", "primary_course_number": "18.014", "course_title": "Calc"})
    )

    def resource(rid, filename, kind, pages):
        (d / "resources" / rid).mkdir(parents=True)
        meta = {"file": f"/courses/x/{filename}", "learning_resource_types": [kind], "parent_uid": "notes-uid"}
        (d / "resources" / rid / "data.json").write_text(json.dumps(meta))
        make_pdf(d / "static_resources" / filename, pages)

    resource("chb", "abc_ChBnotes.pdf", "Lecture Notes", ["Square roots\nRoot two is irrational."])
    resource("cha", "def_ChAnotes.pdf", "Lecture Notes", ["Integers\nInduction.", "Exponents\nPowers."])
    resource("all", "ghi_course_notes.pdf", "Lecture Notes", ["Everything"])
    resource("ps1", "jkl_pset1.pdf", "Problem Sets", ["Problem 1"])
    # Chapter B's second page is a scan: an image with no text layer.
    with pymupdf.open(d / "static_resources" / "abc_ChBnotes.pdf") as doc:
        pixmap = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 8, 8))
        pixmap.clear_with(200)
        doc.new_page().insert_image(pymupdf.Rect(0, 0, 100, 100), pixmap=pixmap)
        doc.saveIncr()
    (d / "pages" / "notes").mkdir(parents=True)
    html = (
        '<p>All notes in one file (<a href="../../resources/all/index.html">PDF - 3MB</a>).</p>'
        "<table><tr><th>CH</th><th>TOPICS</th></tr>"
        '<tr><td>A</td><td>Integers (<a href="../../resources/cha/index.html">PDF</a>)</td></tr>'
        '<tr><td>B</td><td>Square roots (<a href="../../resources/chb/index.html">PDF</a>)</td></tr></table>'
    )
    (d / "pages" / "notes" / "data.json").write_text(json.dumps({"uid": "notes-uid", "content": html}))
    return d
