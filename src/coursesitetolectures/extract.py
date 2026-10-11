"""Extract course files into sections: one per PDF page, slide, heading or code file.

Each section keeps the file and page/slide it came from, so lecture audio can later be synced back
to the materials.
"""

import json
import re
from dataclasses import dataclass

from bs4 import BeautifulSoup

CODE_EXTS = {
    ".py", ".java", ".c", ".cc", ".cpp", ".h", ".hpp", ".js", ".ts", ".m", ".r", ".jl",
    ".go", ".rs", ".scala", ".hs", ".ml", ".rkt", ".scm", ".lisp", ".sql", ".sh",
}  # fmt: skip
DOC_EXTS = {".pdf", ".pptx", ".html", ".htm", ".md", ".markdown", ".txt", ".tex", ".ipynb"}

MAX_TITLE = 80
MD_HEADING = re.compile(r"^#{1,2} +(.+)$", re.MULTILINE)
TEX_SECTION = re.compile(r"^\\(?:sub)?section\*?\{(.+)\}", re.MULTILINE)
PAGE_NUMBER = re.compile(r"^\s*(?:page\s*)?\d+(?:\s*(?:/|of)\s*\d+)?\s*$", re.IGNORECASE)


@dataclass
class Section:
    title: str  # becomes the "## [N] Title" heading and the audio chapter title
    body: str  # Markdown
    source: str  # file path relative to the course directory
    locator: str | None = None  # "page 3", "slide 2", "cell 4", or None for the whole file
    scanned: bool = False  # a page image with no text layer, to transcribe with Claude


def short(text, limit=MAX_TITLE):
    text = " ".join(text.split())
    return text if len(text) <= limit else text[: limit - 1].rstrip() + "…"


def first_line(text):
    """The first line with a few letters in it, as a title."""
    for line in text.splitlines():
        if sum(c.isalpha() for c in line) >= 3:
            return short(line)
    return None


def extract(path, rel):
    """Sections for the file at path; rel is its path relative to the course directory."""
    ext = path.suffix.lower()
    if ext in CODE_EXTS:
        code = path.read_text(errors="replace").strip()
        return [Section(f"Example code: {path.name}", f"```{ext[1:]}\n{code}\n```", rel)]
    if ext == ".pdf":
        return pdf(path, rel)
    if ext == ".pptx":
        return pptx(path, rel)
    if ext in (".html", ".htm"):
        return html(path, rel)
    if ext == ".ipynb":
        return notebook(path, rel)
    text = path.read_text(errors="replace")
    pattern = TEX_SECTION if ext == ".tex" else MD_HEADING
    return split_headings(text, pattern, rel, path.stem)


def pdf(path, rel):
    import pymupdf

    sections = []
    with pymupdf.open(path) as doc:
        for i, page in enumerate(doc, 1):
            lines = [line for line in page.get_text().splitlines() if not PAGE_NUMBER.match(line)]
            text = "\n".join(lines).strip()
            if text:
                sections.append(Section(first_line(text) or f"{path.name}, page {i}", text, rel, f"page {i}"))
            elif page.get_images():
                sections.append(Section(f"{path.name}, page {i}", "", rel, f"page {i}", scanned=True))
    return sections


def pptx(path, rel):
    from pptx import Presentation

    sections = []
    for i, slide in enumerate(Presentation(path).slides, 1):
        title_shape = slide.shapes.title
        title = title_shape.text_frame.text.strip() if title_shape is not None else ""
        title_id = title_shape.shape_id if title_shape is not None else None
        parts = [
            s.text_frame.text.strip()
            for s in slide.shapes
            if s.has_text_frame and s.shape_id != title_id and s.text_frame.text.strip()
        ]
        if slide.has_notes_slide and (notes := slide.notes_slide.notes_text_frame.text.strip()):
            parts.append(f"Speaker notes:\n\n{notes}")  # often the best narration material
        if title or parts:
            sections.append(Section(short(title) or f"{path.name}, slide {i}", "\n\n".join(parts), rel, f"slide {i}"))
    return sections


def html(path, rel):
    soup = BeautifulSoup(path.read_bytes(), "html.parser")
    for tag in soup(["script", "style", "nav", "header", "footer", "noscript", "form"]):
        tag.decompose()
    main = soup.find("main") or soup.find("article") or soup.body or soup
    for h in main.find_all(["h1", "h2"]):
        h.replace_with(f"\n# {h.get_text(' ', strip=True)}\n")
    text = re.sub(r"\n\s*\n+", "\n\n", main.get_text("\n"))
    page_title = soup.title.get_text(strip=True) if soup.title else path.stem
    return split_headings(text, MD_HEADING, rel, page_title)


def notebook(path, rel):
    """One section per Markdown heading cell, with the code cells that follow it."""
    nb = json.loads(path.read_text())
    sections = []
    for i, cell in enumerate(nb.get("cells", []), 1):
        src = cell.get("source", "")
        text = ("".join(src) if isinstance(src, list) else src).strip()
        if not text:
            continue
        if cell.get("cell_type") == "code":
            text = f"```python\n{text}\n```"
        heading = MD_HEADING.match(text) if cell.get("cell_type") == "markdown" else None
        if heading or not sections:
            title = short(heading.group(1)) if heading else path.stem
            sections.append(Section(title, text, rel, f"cell {i}"))
        else:
            sections[-1].body += f"\n\n{text}"
    return sections


def split_headings(text, pattern, rel, default_title):
    """Split text at headings matching pattern; text before the first heading gets default_title."""
    sections, matches = [], list(pattern.finditer(text))
    bounds = [(None, 0)] + [(m, m.start()) for m in matches]
    for (m, start), (_, end) in zip(bounds, bounds[1:] + [(None, len(text))], strict=True):
        body = text[m.end() if m else start : end].strip()
        title = short(m.group(1)) if m else short(default_title)
        if body or m:
            sections.append(Section(title, body, rel, f"section {len(sections) + 1}" if matches else None))
    return sections
